#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the JointAdaSpec defense deck (14 slides, 16:9) as a .pptx.

LEGACY NOTE: this auto-builder produces the earlier 14-slide deck. The current
defense structure is 12 slides (less theory, more results, no standalone
speculative-decoding primer) — see ``papers/pres.md`` and
``papers/claude_design.md``; the authoritative visual deck is built from that
brief. This script is kept as a reproducible legacy artifact and is not the
source of truth for the slide structure.

Content is taken verbatim from ``papers/pres.md`` and the visual system from
``papers/claude_design.md``. Historical large-scale numbers are shown as
artifact-backed thesis-snapshot results, not as post-defense block-v1 reruns. The honest framing
(AR is the fastest method; "2.22×" is relative to vanilla speculative decoding;
joint ≈ cascade is a statistical tie; 7B/1.5B is a null; rerun is required for
future main-branch benchmark claims) is preserved exactly.

Output: ``papers/dist/JointAdaSpec_defense.pptx``. The PDF and HTML renditions
are produced from this file by ``scripts/build_slides.sh`` via LibreOffice.

Run:  .venv/bin/python papers/build_slides.py
"""

import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FIG = os.path.join(ROOT, "reports", "thesis_figs", "png_slides")
OUT = os.path.join(HERE, "dist", "JointAdaSpec_defense.pptx")

# ---- design system (papers/claude_design.md §3) --------------------------------
INDIGO = RGBColor(0x1B, 0x2A, 0x4A)   # headings / bars
GRAPHITE = RGBColor(0x1F, 0x24, 0x30)  # body text
TURQUOISE = RGBColor(0x2B, 0xB3, 0xA3)  # accent / JointAdaSpec
GREEN = RGBColor(0x1E, 0x9E, 0x5A)     # ONLY statistically significant (p<0.05)
GREY = RGBColor(0x8A, 0x8F, 0x98)      # caveats / null
AMBER = RGBColor(0xB5, 0x85, 0x2A)     # honest caveats
CAPTION = RGBColor(0x5A, 0x61, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)     # zebra / panels

# "Calibri" renders as installed Carlito on Linux/LibreOffice (metric-identical,
# full Cyrillic) and as real Calibri in PowerPoint on Windows.
DISPLAY = "Calibri"
MONO = "DejaVu Sans Mono"

EMU_IN = 914400
SW, SH = 13.333, 7.5  # 16:9 inches

# image native sizes (px) for aspect-preserving fit
DIMS = {
    "fig1_pareto.png": (3799, 1509),
    "fig_E_adaptivity_ablation.png": (3938, 1303),
    "fig_bonus_kappa_sweep.png": (2292, 1395),
    "fig_G_acceptance_em.png": (2400, 1245),
    "fig_3win_robustness.png": (2294, 1252),
    "fig_D_advantage_on_B.png": (2390, 1249),
    "repo_qr.png": (468, 468),
}


# ---- low-level helpers ---------------------------------------------------------
def _no_line(shape):
    shape.line.fill.background()


def _solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    _no_line(shape)


def rect(slide, l, t, w, h, color, line=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    _solid(sp, color)
    if line is not None:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def rrect(slide, l, t, w, h, fill, line=None, text=None, size=14, color=None, bold=True):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    _solid(sp, fill)
    if line is not None:
        sp.line.color.rgb = line
        sp.line.width = Pt(1.25)
    sp.shadow.inherit = False
    if text is not None:
        tf = sp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for m in (tf.margin_left, tf.margin_right):
            pass
        tf.margin_left = Inches(0.12); tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.06); tf.margin_bottom = Inches(0.06)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold
        r.font.name = DISPLAY
        r.font.color.rgb = color if color else GRAPHITE
    return sp


def textbox(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    return tf


def add_para(tf, first=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    return p


def run(p, text, size=20, color=GRAPHITE, bold=False, italic=False, name=DISPLAY):
    r = p.add_run(); r.text = text
    f = r.font
    f.size = Pt(size); f.bold = bold; f.italic = italic
    f.name = name; f.color.rgb = color
    return r


def fit_box(path_key, box_w, box_h):
    """Return (w,h) in inches fitting native aspect inside box (inches)."""
    pw, ph = DIMS[path_key]
    ar = pw / ph
    w, h = box_w, box_w / ar
    if h > box_h:
        h, w = box_h, box_h * ar
    return w, h


def picture(slide, key, cx, cy, box_w, box_h, folder=FIG):
    w, h = fit_box(key, box_w, box_h)
    l = cx - w / 2.0
    t = cy - h / 2.0
    slide.shapes.add_picture(os.path.join(folder, key), Inches(l), Inches(t), Inches(w), Inches(h))


# ---- slide scaffold ------------------------------------------------------------
def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    rect(s, 0, 0, SW, SH, WHITE)  # explicit white bg
    return s


def header(slide, title, kicker=None, idx=None):
    rect(slide, 0, 0, 0.22, SH, INDIGO)            # left spine
    rect(slide, 0.7, 1.18, 4.2, 0.05, TURQUOISE)    # accent rule under title
    if kicker:
        tf = textbox(slide, 0.72, 0.34, 11.5, 0.4)
        run(add_para(tf, True), kicker.upper(), size=12, color=TURQUOISE, bold=True)
    tf = textbox(slide, 0.7, 0.52, 12.0, 0.7)
    run(add_para(tf, True), title, size=30, color=INDIGO, bold=True)
    # footer
    tf = textbox(slide, 0.7, SH - 0.42, 9.0, 0.3)
    run(add_para(tf, True), "JointAdaSpec · Козин А.А. · КубГУ ФПМ · 2026", size=10, color=GREY)
    if idx is not None:
        tf = textbox(slide, SW - 1.2, SH - 0.42, 0.9, 0.3)
        p = add_para(tf, True); p.alignment = PP_ALIGN.RIGHT
        run(p, "%d / 14" % idx, size=10, color=GREY)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def bullets(slide, items, left=0.85, top=1.5, width=11.6, size=21, gap=10):
    """items: list of (text, color, bold) or (text,)."""
    tf = textbox(slide, left, top, width, SH - top - 0.6)
    for i, it in enumerate(items):
        text = it[0]
        color = it[1] if len(it) > 1 and it[1] else GRAPHITE
        bold = it[2] if len(it) > 2 else False
        p = add_para(tf, first=(i == 0))
        p.space_after = Pt(gap)
        run(p, "—  ", size=size, color=TURQUOISE, bold=True)
        run(p, text, size=size, color=color, bold=bold)
    return tf


# ---- the 14 slides -------------------------------------------------------------
def slide_title(prs):
    s = new_slide(prs)
    rect(s, 0, 0, SW, 0.22, INDIGO)
    rect(s, 0, SH - 0.22, SW, 0.22, TURQUOISE)
    tf = textbox(s, 1.0, 2.0, 11.33, 2.2, anchor=MSO_ANCHOR.MIDDLE)
    p = add_para(tf, True); p.alignment = PP_ALIGN.CENTER
    run(p, "Оптимизация больших языковых моделей\nметодом адаптивного спекулятивного декодирования", size=32, color=INDIGO, bold=True)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    run(p2, "Метод JointAdaSpec на основе марковского процесса принятия решений", size=20, color=TURQUOISE, bold=True)
    tf2 = textbox(s, 1.0, 4.5, 11.33, 2.0, anchor=MSO_ANCHOR.TOP)
    for i, (txt, c, b, sz) in enumerate([
        ("Выпускная квалификационная работа (магистерская диссертация)", GRAPHITE, True, 18),
        ("Козин Александр Александрович", GRAPHITE, False, 18),
        ("Научный руководитель: Калайдина Г.В.", CAPTION, False, 15),
        ("Кубанский государственный университет · факультет прикладной математики", CAPTION, False, 15),
        ("кафедра математического моделирования · Краснодар, 2026", CAPTION, False, 15),
    ]):
        p = add_para(tf2, first=(i == 0)); p.alignment = PP_ALIGN.CENTER; p.space_after = Pt(6)
        run(p, txt, size=sz, color=c, bold=b)
    notes(s, "Уважаемые члены комиссии, тема моей работы — оптимизация больших языковых моделей "
             "методом адаптивного спекулятивного декодирования. Докладывает Козин Александр.")


def slide_problem(prs):
    s = new_slide(prs)
    header(s, "Актуальность и проблема", "Зачем это нужно", 2)
    bullets(s, [
        ("Стоимость LLM сместилась на инференс — 60–90 % полной стоимости жизненного цикла модели.",),
        ("Авторегрессионная генерация: n последовательных непараллелизуемых проходов; режим, ограниченный памятью.",),
        ("Спекулятивное декодирование (SD) — промышленный стандарт ускорения (vLLM, TensorRT-LLM).",),
        ("Проблема: длина черновика γ и порог верификации T фиксированы и неадаптивны.", INDIGO, True),
    ], top=1.6, size=22, gap=16)
    notes(s, "Основные затраты на LLM — это инференс, а его узкое место — последовательная генерация. "
             "SD ускоряет её, но его ключевые гиперпараметры фиксированы и не подстраиваются под состояние генерации.")


def _flow(slide, items, top, box_w=2.55, box_h=1.0, gap=0.45, fill=LIGHT, line=TURQUOISE, size=14):
    n = len(items)
    total = n * box_w + (n - 1) * gap
    l = (SW - total) / 2.0
    for i, txt in enumerate(items):
        rrect(slide, l, top, box_w, box_h, fill, line=line, text=txt, size=size, color=INDIGO)
        if i < n - 1:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(l + box_w + 0.04),
                                        Inches(top + box_h / 2 - 0.13), Inches(gap - 0.08), Inches(0.26))
            _solid(ar, TURQUOISE); ar.shadow.inherit = False
        l += box_w + gap


def slide_sd_idea(prs):
    s = new_slide(prs)
    header(s, "Спекулятивное декодирование: идея и две оси", "Базовый механизм", 3)
    bullets(s, [
        ("Дешёвая черновая модель предлагает блок токенов; дорогая целевая проверяет его за один проход.",),
        ("Modified rejection sampling гарантирует точное распределение — без потерь качества при T = 1.",),
        ("Две независимые оси управления: длина черновика γ и порог верификации T.", INDIGO, True),
    ], top=1.55, size=21, gap=12)
    _flow(s, ["Черновик q:\nблок γ токенов", "Целевая p:\nпроверка за 1 проход",
              "Принятый префикс\n+ ресемпл"], top=4.35, box_w=3.0, size=15)
    tf = textbox(s, 0.85, 5.7, 11.6, 0.6)
    p = add_para(tf, True); p.alignment = PP_ALIGN.CENTER
    run(p, "Две «ручки» управления:  γ — сколько токенов предлагать   ·   T — насколько строго принимать",
        size=15, color=CAPTION, italic=True)
    notes(s, "Дешёвый черновик угадывает несколько токенов, дорогая модель проверяет их разом. "
             "Управлять можно двумя ручками — длиной черновика и порогом приёмки.")


def slide_taxonomy(prs):
    s = new_slide(prs)
    header(s, "Сравнение методов и исследовательский пробел", "Место работы в литературе", 4)
    cols = ["", "Длина: фиксир.", "Длина: эвристич.", "Длина: обучаемая"]
    data = [
        ["Вериф.: строгая", "классич. SD", "DISCO", "SpecDec++, SVIP"],
        ["Вериф.: ослабл. фикс.", "AutoJudge, Fuzzy SD", "—", "—"],
        ["Вериф.: ослабл. адаптивн.", "MARS", "—", "JointAdaSpec ★"],
    ]
    L, T, W, H = 1.1, 1.75, 11.1, 3.6
    tbl = s.shapes.add_table(4, 4, Inches(L), Inches(T), Inches(W), Inches(H)).table
    tbl.columns[0].width = Inches(3.0)
    for c in range(1, 4):
        tbl.columns[c].width = Inches((W - 3.0) / 3)
    for c, txt in enumerate(cols):
        cell = tbl.cell(0, c); cell.text = ""
        cell.fill.solid(); cell.fill.fore_color.rgb = INDIGO
        p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = txt; r.font.size = Pt(13); r.font.bold = True
        r.font.color.rgb = WHITE; r.font.name = DISPLAY
    for ri, row in enumerate(data, start=1):
        for ci, txt in enumerate(row):
            cell = tbl.cell(ri, ci)
            highlight = (txt == "JointAdaSpec ★")
            cell.fill.solid()
            cell.fill.fore_color.rgb = TURQUOISE if highlight else (LIGHT if ci == 0 else WHITE)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER if ci else PP_ALIGN.LEFT
            r = p.add_run(); r.text = txt
            r.font.size = Pt(13); r.font.bold = (ci == 0 or highlight)
            r.font.color.rgb = WHITE if highlight else (INDIGO if ci == 0 else GRAPHITE)
            r.font.name = DISPLAY
    tf = textbox(s, 1.1, 5.7, 11.1, 0.8)
    p = add_para(tf, True)
    run(p, "Существующие методы адаптируют ровно одну ось. Пустая клетка ", size=18, color=GRAPHITE)
    run(p, "«обучаемая длина × адаптивная верификация»", size=18, color=INDIGO, bold=True)
    run(p, " — её закрывает JointAdaSpec.", size=18, color=GRAPHITE)
    notes(s, "Почти все методы адаптируют одну ось, а вторую держат фиксированной. "
             "Совместного адаптивного управления обеими осями в литературе не было — эту клетку закрывает работа.")


def slide_goal(prs):
    s = new_slide(prs)
    header(s, "Цель и задачи", "Постановка", 5)
    panel = rrect(s, 0.85, 1.55, 11.6, 1.15, LIGHT, line=TURQUOISE)
    tf = panel.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    run(p, "Цель: ", size=20, color=INDIGO, bold=True)
    run(p, "метод адаптивного спекулятивного декодирования, совместно управляющий парой "
          "(длина черновика, порог верификации) как задачей оптимального управления.", size=20, color=GRAPHITE)
    items = [
        "Анализ методов ускорения инференса и точная формулировка исследовательского пробела.",
        "MDP-модель совместного управления и её теоретический анализ.",
        "Воспроизводимый программный комплекс.",
        "Экспериментальная оценка и установление границ применимости.",
    ]
    tf2 = textbox(s, 0.95, 3.05, 11.5, 3.4)
    for i, it in enumerate(items):
        p = add_para(tf2, first=(i == 0)); p.space_after = Pt(14)
        run(p, "%d.  " % (i + 1), size=22, color=TURQUOISE, bold=True)
        run(p, it, size=21, color=GRAPHITE)
    notes(s, "Цель — формализовать совместное управление двумя осями. "
             "Для этого поставлены четыре задачи: от анализа литературы до воспроизводимого эксперимента.")


def slide_method(prs):
    s = new_slide(prs)
    header(s, "Метод JointAdaSpec", "Что сделано", 6)
    bullets(s, [
        ("Каждый цикл декодирования = выбор действия в дискретном MDP (S, A, P, r, λ).",),
        ("Состояние s = (энтропия черновика H, дивергенция K = KL(q‖p), позиция k).",),
        ("Действие a ∈ {continue} ∪ {verify@T}; |S| = 3600, |A| = 9.",),
        ("Награда r = η − κ·D: скорость минус штраф за качество; решение — value iteration, λ = 0,99.",),
        ("Политика — таблица ≈ 2 КБ, поиск O(1) на шаг, без обучения нейросетей.", INDIGO, True),
    ], top=1.5, size=20, gap=9)
    _flow(s, ["Состояние\n(H, K, k)", "Политика-таблица\n(value iteration)", "Действие\ncontinue/verify@T"],
          top=5.55, box_w=3.0, box_h=0.95, size=15)
    notes(s, "Каждый шаг — выбор действия в марковском процессе: состояние — признаки генерации, "
             "действие — пара (длина, порог). Процесс маленький, 3600 состояний, решается точно за секунды.")


def slide_theory(prs):
    s = new_slide(prs)
    header(s, "Теоретические результаты", "Шесть результатов", 7)
    cards = [
        ("A", "Выборочная сложность оценщика переходов", False),
        ("B", "исправленная reward-shaping формулировка", False),
        ("C", "Оценка субоптимальности каскада через массу нарушения μ*_J(B)", False),
        ("D", "Точный разрыв ценности совместной и каскадной политик", True),
        ("2.3", "Слабое доминирование совместной политики над каскадной", False),
        ("2.4", "Скаляризация Парето-фронта по коэффициенту κ", False),
    ]
    cw, ch, gx, gy = 5.55, 1.35, 0.4, 0.35
    L0, T0 = 0.85, 1.7
    for i, (tag, desc, hl) in enumerate(cards):
        r, c = divmod(i, 2)
        l = L0 + c * (cw + gx); t = T0 + r * (ch + gy)
        card = rrect(s, l, t, cw, ch, INDIGO if hl else LIGHT,
                     line=TURQUOISE if hl else GREY)
        tf = card.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.18); tf.margin_right = Inches(0.15)
        p = tf.paragraphs[0]
        run(p, "Теорема %s. " % tag, size=16, color=TURQUOISE if hl else INDIGO, bold=True)
        run(p, desc, size=15, color=WHITE if hl else GRAPHITE, bold=False)
    notes(s, "Получен аппарат из шести результатов. Ключевой — теорема D: она точно выражает разрыв между "
             "совместной и каскадной политиками и объясняет, почему на практике они совпадают.")


def slide_impl(prs):
    s = new_slide(prs)
    header(s, "Программная реализация и воспроизводимость", "Инженерная часть", 8)
    _flow(s, ["01\nСбор трейсов", "02\nРешение MDP", "03\nБенчмарк", "04\nПроверка условий"],
          top=1.7, box_w=2.5, box_h=1.0, gap=0.4, size=15)
    bullets(s, [
        ("Стек: Python 3.12, PyTorch, scipy.sparse, Hydra; полный pytest-набор: 94 теста на v1-defense.",),
        ("Воспроизводимость: манифесты (git SHA, сиды [42, 43, 44], SHA256), детерминированные прогоны.",),
        ("Совместимость с инференс-фреймворками без переобучения базовых моделей.",),
    ], top=3.25, size=21, gap=14)
    notes(s, "Метод реализован как воспроизводимый конвейер: каждый прогон детерминирован и "
             "сопровождается манифестом; на defense snapshot полный pytest-набор прошёл как 94 теста.")


def _results_table(slide, L, T, W, H):
    head = ["Метод", "EM, %", "ток/с", "Δ EM к target", "p"]
    rows = [
        ("target_only", "52,93", "14,45", "—", "—", None),
        ("speculative", "53,27", "4,88", "+0,33", "0,877", None),
        ("cascade", "57,13", "10,29", "+4,20", "0,015", None),
        ("jointadaspec", "57,00", "10,84", "+4,07", "0,020", TURQUOISE),
    ]
    tbl = slide.shapes.add_table(5, 5, Inches(L), Inches(T), Inches(W), Inches(H)).table
    widths = [2.6, 1.3, 1.3, 2.1, 1.2]
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    for c, txt in enumerate(head):
        cell = tbl.cell(0, c); cell.fill.solid(); cell.fill.fore_color.rgb = INDIGO
        cell.text = ""
        p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        r = p.add_run(); r.text = txt; r.font.size = Pt(14); r.font.bold = True
        r.font.color.rgb = WHITE; r.font.name = DISPLAY
    for ri, row in enumerate(rows, start=1):
        hl = row[5]
        for ci, txt in enumerate(row[:5]):
            cell = tbl.cell(ri, ci); cell.fill.solid()
            cell.fill.fore_color.rgb = hl if hl else (LIGHT if ri % 2 else WHITE)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
            r = p.add_run(); r.text = txt
            r.font.size = Pt(14); r.font.bold = bool(hl) or ci == 0
            r.font.color.rgb = WHITE if hl else GRAPHITE; r.font.name = DISPLAY


def slide_results_main(prs):
    s = new_slide(prs)
    header(s, "Thesis snapshot: пара 14B / 0.5B", "artifact-backed result · GSM8K · n = 1500", 9)
    _results_table(s, 0.85, 1.55, 8.5, 2.5)
    # big-number callout (green: statistically significant)
    panel = rrect(s, 0.85, 4.35, 5.6, 2.4, LIGHT, line=GREEN)
    tf = panel.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, "+4,07 п.п.", size=50, color=GREEN, bold=True)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    run(p2, "snapshot EM к target (p = 0,0205)", size=15, color=GRAPHITE)
    p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
    run(p3, "2,22× ", size=26, color=INDIGO, bold=True)
    run(p3, "пропускной способности (к ванильному SD)", size=14, color=CAPTION)
    p4 = tf.add_paragraph(); p4.alignment = PP_ALIGN.CENTER
    run(p4, "artifact-backed thesis result; fresh claims — в main", size=12, color=AMBER, bold=True)
    picture(s, "fig1_pareto.png", cx=9.85, cy=4.35, box_w=6.6, box_h=2.5)
    tf2 = textbox(s, 6.7, 5.7, 6.3, 0.8)
    p = add_para(tf2, True); p.alignment = PP_ALIGN.CENTER
    run(p, "Парето-фронт: прямая генерация (target_only) — самая быстрая точка.", size=13, color=CAPTION, italic=True)
    notes(s, "Это зафиксированный artifact-backed результат snapshot ВКР. Новые post-defense block-v1 claims "
             "делаются отдельно в main, чтобы не смешивать поколения артефактов.")


def slide_ablation(prs):
    s = new_slide(prs)
    header(s, "Адаптивность нетривиальна: эксперимент E", "thesis-snapshot ablation · n = 300", 10)
    picture(s, "fig_E_adaptivity_ablation.png", cx=SW / 2, cy=4.0, box_w=11.4, box_h=4.0)
    panel = rrect(s, 2.4, 6.15, 8.5, 0.95, LIGHT, line=GREEN)
    tf = panel.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, "JointAdaSpec точнее любого фиксированного T на ", size=16, color=GRAPHITE)
    run(p, "+4…+8 п.п.", size=16, color=GREEN, bold=True)
    run(p, " EM и быстрее ", size=16, color=GRAPHITE)
    run(p, "≈ 3,7×", size=16, color=INDIGO, bold=True)
    run(p, " (11,12 против ≈ 3,0 ток/с).", size=16, color=GRAPHITE)
    tf2 = textbox(s, 2.4, 7.03, 8.5, 0.35)
    pp = add_para(tf2, True); pp.alignment = PP_ALIGN.CENTER
    run(pp, "Ablation evidence для snapshot ВКР; новые block-v1 claims требуют rerun в main.", size=10, color=AMBER, bold=True)
    notes(s, "Нужна ли адаптивность, если порог можно подобрать? Thesis-snapshot ablation показывает "
             "мотивацию, а новые post-defense claims выносятся в main.")


def slide_honest(prs):
    s = new_slide(prs)
    header(s, "Честная картина результатов", "Границы применимости", 11)
    bullets(s, [
        ("joint ≈ cascade — статистическая ничья (Δ = −0,13 %, p = 0,96), ровно как предсказывает теорема D.", AMBER, False),
        ("Прямая генерация (AR) — самый быстрый метод; «2,22×» — относительно ванильного SD, не AR.", GREY, False),
        ("Пара 7B / 1.5B (4,7×) — нуль-результат, устойчивый по триангуляции на трёх независимых окнах.", AMBER, False),
        ("Прирост качества немонотонен: +7,4 п.п. при умеренной доле принятий, −2,6 п.п. при высокой.", GREY, False),
    ], left=0.85, top=1.55, width=7.3, size=18, gap=14)
    picture(s, "fig_bonus_kappa_sweep.png", cx=10.4, cy=3.4, box_w=4.6, box_h=3.2)
    tf = textbox(s, 8.2, 5.15, 4.5, 0.6)
    p = add_para(tf, True); p.alignment = PP_ALIGN.CENTER
    run(p, "κ-развёртка: joint ≈ cascade на всём диапазоне.", size=12, color=CAPTION, italic=True)
    tf2 = textbox(s, 0.85, 6.5, 11.6, 0.7)
    p = add_para(tf2, True)
    run(p, "Это не слабость изложения, а строгая характеризация области применимости.", size=16, color=INDIGO, bold=True)
    notes(s, "Честно обозначаю границы: совместная политика не превосходит каскадную (предсказание теории), "
             "прямая генерация быстрее всех, на паре с низким отношением мощностей эффекта нет.")


def slide_conclusion(prs):
    s = new_slide(prs)
    header(s, "Заключение и вклад", "Итог работы", 12)
    cards = [
        ("Вклад 1", "Единый MDP-фреймворк совместного управления (длина, порог) — одна из первых табличных постановок в литературе."),
        ("Вклад 2", "Теоретический аппарат A / B / C / D / 2.3 / 2.4 с доказательствами."),
        ("Вклад 3", "Честная эмпирическая характеризация: где адаптивность помогает, где нет и почему joint сводится к каскаду."),
    ]
    T0 = 1.65
    for i, (tag, desc) in enumerate(cards):
        t = T0 + i * 1.25
        rrect(s, 0.85, t, 1.7, 1.05, TURQUOISE, text=tag, size=17, color=WHITE)
        panel = rrect(s, 2.7, t, 9.75, 1.05, LIGHT, line=GREY)
        tf = panel.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        run(tf.paragraphs[0], desc, size=18, color=GRAPHITE)
    tf = textbox(s, 0.85, 5.6, 11.6, 0.8)
    p = add_para(tf, True)
    run(p, "Направления: ", size=16, color=INDIGO, bold=True)
    run(p, "древовидное обобщение, распределительные признаки состояния, GPU-резидентная политика.", size=16, color=GRAPHITE)
    notes(s, "Итог — не отдельный рекорд, а единый фреймворк с доказанными свойствами и честной картиной того, "
             "когда совместное адаптивное управление действительно полезно.")


def slide_defense(prs):
    s = new_slide(prs)
    header(s, "Положения, выносимые на защиту", "Тезисы защиты", 13)
    items = [
        "Совместное управление (длина, порог) формализуемо как конечный MDP (≈ 3600 состояний), решаемый точно методом value iteration без обучения нейросетевых модулей.",
        "Доказаны выборочная сложность (A), исправленная reward-shaping формулировка (B), оценка и точный разрыв совместной и каскадной политик (C, D), слабое доминирование (2.3) и Парето-скаляризация (2.4).",
        "Реализован block-v1 pipeline с семантической валидацией: разные поколения policies не смешиваются неявно.",
        "Зафиксированные прогоны показывают потенциал; fresh block-v1 claims выносятся в post-defense main. Установлены границы применимости: joint ≈ cascade и нуль на низком отношении мощностей.",
    ]
    tf = textbox(s, 0.95, 1.6, 11.6, 5.3)
    for i, it in enumerate(items):
        p = add_para(tf, first=(i == 0)); p.space_after = Pt(13)
        run(p, "%d.  " % (i + 1), size=19, color=TURQUOISE, bold=True)
        run(p, it, size=18, color=GRAPHITE)
    notes(s, "На защиту выносятся четыре положения: формализуемость совместного управления как конечного MDP; "
             "доказанные теоретические свойства; исправленная реализация и защита артефактов; границы применимости.")


def slide_thanks(prs):
    s = new_slide(prs)
    rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, SW, 0.22, INDIGO)
    rect(s, 0, SH - 0.22, SW, 0.22, TURQUOISE)
    tf = textbox(s, 0.9, 1.4, 7.6, 4.6, anchor=MSO_ANCHOR.MIDDLE)
    p = add_para(tf, True)
    run(p, "Спасибо за внимание", size=40, color=INDIGO, bold=True)
    p2 = tf.add_paragraph(); p2.space_before = Pt(18)
    run(p2, "Готов ответить на ваши вопросы.", size=20, color=GRAPHITE)
    p3 = tf.add_paragraph(); p3.space_before = Pt(24)
    run(p3, "Репозиторий: код, 94 pytest-теста, манифесты воспроизводимости,", size=15, color=CAPTION)
    p4 = tf.add_paragraph()
    run(p4, "текст ВКР, deck, audit/Q&A и артефакты.", size=15, color=CAPTION)
    p5 = tf.add_paragraph(); p5.space_before = Pt(10)
    run(p5, "github.com/levvius/adaptive-speculative-decoding/tree/v1-defense", size=13, color=TURQUOISE, bold=True, name=MONO)
    qr = os.path.join(HERE, "repo_qr.png")
    if os.path.exists(qr):
        s.shapes.add_picture(qr, Inches(9.0), Inches(2.35), Inches(2.8), Inches(2.8))
        tf2 = textbox(s, 9.0, 5.2, 2.8, 0.4)
        pp = add_para(tf2, True); pp.alignment = PP_ALIGN.CENTER
        run(pp, "→ репозиторий", size=12, color=CAPTION)
    notes(s, "Код, тесты, манифесты, deck, audit/Q&A и текст работы открыты — QR ведёт на immutable snapshot v1-defense. "
             "Спасибо за внимание, готов ответить на вопросы.")


def main():
    prs = Presentation()
    prs.slide_width = Emu(int(SW * EMU_IN))
    prs.slide_height = Emu(int(SH * EMU_IN))
    slide_title(prs)
    slide_problem(prs)
    slide_sd_idea(prs)
    slide_taxonomy(prs)
    slide_goal(prs)
    slide_method(prs)
    slide_theory(prs)
    slide_impl(prs)
    slide_results_main(prs)
    slide_ablation(prs)
    slide_honest(prs)
    slide_conclusion(prs)
    slide_defense(prs)
    slide_thanks(prs)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print("wrote %s (%d slides)" % (OUT, len(prs.slides._sldIdLst)))


if __name__ == "__main__":
    main()
